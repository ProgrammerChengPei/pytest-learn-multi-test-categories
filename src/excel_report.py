"""
Excel Report Generator for Pytest / Pytest Excel 报告生成器
"""
import json
import time
from pathlib import Path
from typing import Dict, Any, List
from datetime import datetime


class ExcelReportGenerator:
    """Excel report generator class / Excel报告生成器类"""

    def __init__(self, output_dir: Path):
        self.output_dir = output_dir
        self.output_dir.mkdir(exist_ok=True)
        self.test_results: List[Dict[str, Any]] = []
        self.start_time: float = 0
        self.end_time: float = 0

    def start(self):
        """Start timing / 开始计时"""
        self.start_time = time.time()

    def stop(self):
        """Stop timing / 停止计时"""
        self.end_time = time.time()

    def add_test_result(self, result: Dict[str, Any]):
        """
        Add test result / 添加测试结果

        Args:
            result: Test result dictionary / 测试结果字典
        """
        self.test_results.append(result)

    def generate_report(self, filename: str = "test_report.xlsx") -> str:
        """
        Generate Excel report / 生成Excel报告

        Args:
            filename: Output filename / 输出文件名

        Returns:
            Path to generated report / 生成的报告路径
        """
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            print("Warning: openpyxl not installed. Generating CSV instead...")
            return self.generate_csv_report(filename.replace('.xlsx', '.csv'))

        wb = Workbook()

        # 删除默认工作表
        if 'Sheet' in wb.sheetnames:
            del wb['Sheet']

        # 创建各个工作表
        self._create_summary_sheet(wb)
        self._create_flash_tests_sheet(wb)
        self._create_interface_tests_sheet(wb)
        self._create_function_tests_sheet(wb)
        self._create_performance_tests_sheet(wb)

        # 保存文件
        output_path = self.output_dir / filename
        wb.save(str(output_path))
        print(f"✓ Excel report generated: {output_path}")
        return str(output_path)

    def _create_summary_sheet(self, wb: 'Workbook'):
        """Create summary sheet / 创建汇总工作表"""
        ws = wb.create_sheet("测试汇总 / Summary", 0)

        # 定义样式
        title_font = Font(name='微软雅黑', size=14, bold=True)
        header_font = Font(name='微软雅黑', size=11, bold=True)
        cell_font = Font(name='微软雅黑', size=10)
        title_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # 标题
        ws.merge_cells('A1:D1')
        title_cell = ws['A1']
        title_cell.value = "Pytest 测试报告 / Pytest Test Report"
        title_cell.font = title_font
        title_cell.fill = title_fill
        title_cell.alignment = Alignment(horizontal='center', vertical='center')

        # 测试信息
        test_info = [
            ["报告生成时间 / Report Time", datetime.now().strftime('%Y-%m-%d %H:%M:%S')],
            ["测试总耗时 / Total Duration", f"{(self.end_time - self.start_time):.2f} 秒 / seconds"],
            ["测试总数 / Total Tests", len(self.test_results)],
            ["通过数 / Passed", len([r for r in self.test_results if r.get('status') == 'passed'])],
            ["失败数 / Failed", len([r for r in self.test_results if r.get('status') == 'failed'])],
            ["跳过数 / Skipped", len([r for r in self.test_results if r.get('status') == 'skipped'])],
        ]

        row = 3
        for label, value in test_info:
            ws[f'A{row}'] = label
            ws[f'B{row}'] = value
            ws[f'A{row}'].font = header_font
            ws[f'B{row}'].font = cell_font
            ws[f'A{row}'].fill = header_fill
            row += 1

        # 按测试类型统计
        ws['A10'] = "按测试类型统计 / Statistics by Type"
        ws['A10'].font = title_font
        ws['A10'].fill = title_fill
        ws.merge_cells('A10:D10')

        type_headers = ["测试类型 / Type", "总数 / Total", "通过 / Passed", "失败 / Failed"]
        for col, header in enumerate(type_headers, 1):
            cell = ws.cell(row=11, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = border

        type_stats = self._calculate_type_statistics()
        row = 12
        for type_name, stats in type_stats.items():
            ws.cell(row=row, column=1, value=type_name).border = border
            ws.cell(row=row, column=2, value=stats['total']).border = border
            ws.cell(row=row, column=3, value=stats['passed']).border = border
            ws.cell(row=row, column=4, value=stats['failed']).border = border
            row += 1

        # 设置列宽
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 12

    def _create_flash_tests_sheet(self, wb: 'Workbook'):
        """Create flash tests sheet / 创建刷写测试工作表"""
        ws = wb.create_sheet("刷写测试 / Flash")

        # 表头
        headers = ["测试名称 / Test Name", "状态 / Status", "耗时 / Duration (s)", "消息 / Message"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(name='微软雅黑', size=11, bold=True)
            cell.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')

        # 填充数据
        flash_tests = [r for r in self.test_results if r.get('category') == 'flash']
        row = 2
        for test in flash_tests:
            ws.cell(row=row, column=1, value=test.get('name', ''))
            status_cell = ws.cell(row=row, column=2, value=test.get('status', ''))
            ws.cell(row=row, column=3, value=f"{test.get('duration', 0):.3f}")
            ws.cell(row=row, column=4, value=test.get('message', ''))

            # 根据状态设置背景色
            if status_cell.value == 'passed':
                status_cell.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
            elif status_cell.value == 'failed':
                status_cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
            row += 1

        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 50

    def _create_interface_tests_sheet(self, wb: 'Workbook'):
        """Create interface tests sheet / 创建接口测试工作表"""
        ws = wb.create_sheet("接口测试 / Interface")

        headers = ["测试名称 / Test Name", "状态 / Status", "耗时 / Duration (s)", "消息 / Message"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(name='微软雅黑', size=11, bold=True)
            cell.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')

        interface_tests = [r for r in self.test_results if r.get('category') == 'interface']
        row = 2
        for test in interface_tests:
            ws.cell(row=row, column=1, value=test.get('name', ''))
            status_cell = ws.cell(row=row, column=2, value=test.get('status', ''))
            ws.cell(row=row, column=3, value=f"{test.get('duration', 0):.3f}")
            ws.cell(row=row, column=4, value=test.get('message', ''))

            if status_cell.value == 'passed':
                status_cell.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
            elif status_cell.value == 'failed':
                status_cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
            row += 1

        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 50

    def _create_function_tests_sheet(self, wb: 'Workbook'):
        """Create function tests sheet / 创建功能测试工作表"""
        ws = wb.create_sheet("功能测试 / Function")

        headers = ["测试名称 / Test Name", "状态 / Status", "耗时 / Duration (s)", "消息 / Message"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(name='微软雅黑', size=11, bold=True)
            cell.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')

        function_tests = [r for r in self.test_results if r.get('category') == 'function']
        row = 2
        for test in function_tests:
            ws.cell(row=row, column=1, value=test.get('name', ''))
            status_cell = ws.cell(row=row, column=2, value=test.get('status', ''))
            ws.cell(row=row, column=3, value=f"{test.get('duration', 0):.3f}")
            ws.cell(row=row, column=4, value=test.get('message', ''))

            if status_cell.value == 'passed':
                status_cell.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
            elif status_cell.value == 'failed':
                status_cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
            row += 1

        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 50

    def _create_performance_tests_sheet(self, wb: 'Workbook'):
        """Create performance tests sheet / 创建性能测试工作表"""
        ws = wb.create_sheet("性能测试 / Performance")

        headers = ["测试名称 / Test Name", "状态 / Status", "指标 / Metric", "值 / Value"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(name='微软雅黑', size=11, bold=True)
            cell.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')

        performance_tests = [r for r in self.test_results if r.get('category') == 'performance']
        row = 2
        for test in performance_tests:
            ws.cell(row=row, column=1, value=test.get('name', ''))
            status_cell = ws.cell(row=row, column=2, value=test.get('status', ''))
            ws.cell(row=row, column=3, value=test.get('metric', ''))
            ws.cell(row=row, column=4, value=str(test.get('value', '')))

            if status_cell.value == 'passed':
                status_cell.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
            elif status_cell.value == 'failed':
                status_cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
            row += 1

        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 25

    def _calculate_type_statistics(self) -> Dict[str, Dict[str, int]]:
        """Calculate statistics by test type / 按测试类型计算统计"""
        stats = {
            '刷写 / Flash': {'total': 0, 'passed': 0, 'failed': 0},
            '接口 / Interface': {'total': 0, 'passed': 0, 'failed': 0},
            '功能 / Function': {'total': 0, 'passed': 0, 'failed': 0},
            '性能 / Performance': {'total': 0, 'passed': 0, 'failed': 0}
        }

        for result in self.test_results:
            category = result.get('category', '')
            status = result.get('status', '')

            if category in stats:
                stats[category]['total'] += 1
                if status == 'passed':
                    stats[category]['passed'] += 1
                elif status == 'failed':
                    stats[category]['failed'] += 1

        return stats

    def generate_csv_report(self, filename: str) -> str:
        """
        Generate CSV report as fallback / 生成CSV报告作为后备方案

        Args:
            filename: Output filename / 输出文件名

        Returns:
            Path to generated report / 生成的报告路径
        """
        import csv

        output_path = self.output_dir / filename
        with open(output_path, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)

            # 汇总信息
            writer.writerow(['测试汇总 / Summary'])
            writer.writerow([])
            writer.writerow(['报告生成时间 / Report Time', datetime.now().strftime('%Y-%m-%d %H:%M:%S')])
            writer.writerow(['测试总耗时 / Total Duration', f'{(self.end_time - self.start_time):.2f} 秒'])
            writer.writerow(['测试总数 / Total Tests', len(self.test_results)])
            writer.writerow([])

            # 测试结果
            writer.writerow(['测试名称 / Test Name', '类型 / Type', '状态 / Status', '耗时 / Duration (s)', '消息 / Message'])
            for result in self.test_results:
                writer.writerow([
                    result.get('name', ''),
                    result.get('category', ''),
                    result.get('status', ''),
                    f"{result.get('duration', 0):.3f}",
                    result.get('message', '')
                ])

        print(f"✓ CSV report generated: {output_path}")
        return str(output_path)


def generate_excel_report_from_artifacts(artifacts_dir: Path) -> str:
    """
    Generate Excel report from test artifacts / 从测试产物生成Excel报告

    Args:
        artifacts_dir: Directory containing test artifacts / 包含测试产物的目录

    Returns:
        Path to generated report / 生成的报告路径
    """
    generator = ExcelReportGenerator(artifacts_dir)

    # Load flash test result
    flash_result_file = artifacts_dir / "flash_test_result.json"
    if flash_result_file.exists():
        with open(flash_result_file, 'r', encoding='utf-8') as f:
            flash_data = json.load(f)
            generator.add_test_result({
                'name': 'Flash Test',
                'category': 'flash',
                'status': flash_data.get('status', 'skipped'),
                'duration': 0,
                'message': flash_data.get('message', 'Flash test completed')
            })

    # Load interface test result
    interface_result_file = artifacts_dir / "interface_test_result.json"
    if interface_result_file.exists():
        with open(interface_result_file, 'r', encoding='utf-8') as f:
            interface_data = json.load(f)
            generator.add_test_result({
                'name': 'Interface Test',
                'category': 'interface',
                'status': interface_data.get('status', 'skipped'),
                'duration': 0,
                'message': interface_data.get('message', 'Interface test completed')
            })

    # Load function test result
    function_result_file = artifacts_dir / "function_test_result.json"
    if function_result_file.exists():
        with open(function_result_file, 'r', encoding='utf-8') as f:
            function_data = json.load(f)
            generator.add_test_result({
                'name': 'Function Test',
                'category': 'function',
                'status': function_data.get('status', 'skipped'),
                'duration': 0,
                'message': function_data.get('message', 'Function test completed')
            })

    # Load performance test result
    performance_result_file = artifacts_dir / "performance_test_result.json"
    if performance_result_file.exists():
        with open(performance_result_file, 'r', encoding='utf-8') as f:
            performance_data = json.load(f)
            generator.add_test_result({
                'name': 'Performance Test',
                'category': 'performance',
                'status': performance_data.get('status', 'skipped'),
                'duration': 0,
                'message': performance_data.get('message', 'Performance test completed')
            })

    return generator.generate_report()
